#!/usr/bin/env python3
"""
Turns a FILLED template + the audio files into an upload package for one language.

  python3 tools/build_language.py --template ko_template.xlsx --audio ko_audio.zip
  python3 tools/build_language.py --template ko_template.xlsx            # text only, audio comes later

What it does (and reports, briefly):
  1. reads the spreadsheet, checks every cell, writes lang/<code>.js
  2. finds audio named <code>_intro.mp3, <code>_s1..s6.mp3, <code>_q01..q15.mp3 (rejects other languages' prefixes),
     converts to MP3 mono 96 kbps where needed, and checks each clip's length against its text (catches swapped files)
  3. makes sure flags/<code>.svg exists, adds the language to languages.js and bumps ASSET_VERSION
  4. runs the checker and builds dist/<code>-update.zip: exactly the files to upload to GitHub

Needs: pip install openpyxl     and ffmpeg/ffprobe on PATH.
"""
import argparse, datetime, json, os, re, shutil, statistics, subprocess, sys, tempfile, zipfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lang_schema as S
from openpyxl import load_workbook

CODE_RE = re.compile(r"^[a-z]{2,3}(-[a-z]{2,4})?$"); TAG_RE = re.compile(r"^[a-z]{2,3}(-[A-Za-z0-9]{2,8})*$")
AUDIO_RE = re.compile(r"^(?P<prefix>.+?)_(?P<slot>intro|s[1-6]|q(?:0[1-9]|1[0-5]))\.mp3$", re.I)
SLOTS = ["intro"] + [f"s{i}" for i in range(1, 7)] + [f"q{i:02d}" for i in range(1, 16)]
errors, warnings, notes = [], [], []
def err(m): errors.append(m)
def warn(m): warnings.append(m)

# ------------------------------------------------------------------ 1. spreadsheet
def cell_text(v):
    if v is None: return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None

def read_template(path):
    wb = load_workbook(path, data_only=True)
    settings = {}
    for r in range(2, wb["Settings"].max_row + 1):
        k = wb["Settings"].cell(r, 1).value
        if k: settings[k] = cell_text(wb["Settings"].cell(r, 3).value)
    tr, en = {}, {}
    for sheet in ("Page 2", "Buttons & labels", "Sections"):
        ws = wb[sheet]
        for r in range(2, ws.max_row + 1):
            fid = ws.cell(r, 1).value
            if fid: en[fid], tr[fid] = cell_text(ws.cell(r, 3).value), cell_text(ws.cell(r, 4).value)
    ws = wb["Questions"]
    for r in range(2, ws.max_row + 1):
        qid = ws.cell(r, 1).value
        if qid:
            en[qid + "_t"], en[qid + "_q"] = cell_text(ws.cell(r, 3).value), cell_text(ws.cell(r, 4).value)
            tr[qid + "_t"], tr[qid + "_q"] = cell_text(ws.cell(r, 5).value), cell_text(ws.cell(r, 6).value)
    return settings, tr, en

def clean(fid, text):
    if text is None: return None
    if fid.startswith("how") or fid.endswith("_t") or (fid.startswith("q") and fid.endswith("_q")):
        text = re.sub(r"^\s*\d+\s*[\.\)]\s*", "", text)               # a '1.' the app already adds
    if fid.endswith("_label"):
        text = re.sub(r"^[^\w(（]+", "", text); text = re.sub(r"[\s—–\-:：]+$", "", text)   # leading emoji, trailing dash
    return text.strip() or None

def build_pack(code, settings, tr, en_tpl, repo):
    base = S.load_pack(os.path.join(repo, "lang", "en.js")); pack = json.loads(json.dumps(base))
    # --- template English must still match the app's English (otherwise rows may have shifted)
    drift = []
    for sh, fid, _, path, _ in S.FIELDS:
        if en_tpl.get(fid) != re.sub(r"\s+", " ", S.get_path(base, path)).strip(): drift.append(fid)
    for i, sid in enumerate(S.SECTION_IDS):
        if en_tpl.get(sid) != base["sections"][i]: drift.append(sid)
    for i, qid in enumerate(S.QUESTION_IDS):
        if en_tpl.get(qid + "_t") != base["questions"][i]["t"] or en_tpl.get(qid + "_q") != base["questions"][i]["q"]: drift.append(qid)
    if drift: err(f"template's English no longer matches the app ({len(drift)} cells, e.g. {', '.join(drift[:4])}). Regenerate the template with tools/make_template.py and copy the translations across.")
    missing, same = [], []
    def put(fid, path, kind="plain"):
        val = clean(fid, tr.get(fid)); english = S.get_path(base, path)
        if val is None: missing.append(fid); return
        if val == english and code != "en": same.append(fid)
        S.set_path(pack, path, val)
    for sh, fid, _, path, keep in S.FIELDS:
        put(fid, path)
        if keep and "{n}" not in S.get_path(pack, path): err(f"'{fid}' must contain {{n}} (it is where the app puts the number)")
    for i, sid in enumerate(S.SECTION_IDS): put(sid, ("sections", i))
    for i, qid in enumerate(S.QUESTION_IDS):
        t, q = clean(qid + "_t", tr.get(qid + "_t")), clean(qid + "_q", tr.get(qid + "_q"))
        if not t or not q: missing.append(qid); continue
        pack["questions"][i] = {"t": t, "q": q}
    # --- settings
    for k in ("code", "name", "gloss", "tag", "dir"):
        if not settings.get(k): err(f"Settings: '{k}' is empty")
    pack.update({"code": code, "name": settings.get("name"), "gloss": settings.get("gloss"), "tag": settings.get("tag"), "dir": settings.get("dir")})
    if settings.get("dir") not in ("ltr", "rtl"): err("Settings: direction must be ltr or rtl")
    if settings.get("tag") and not TAG_RE.match(settings["tag"]): err(f"Settings: tag '{settings['tag']}' looks wrong (expected e.g. fr-FR)")
    reg = {l["code"]: l for l in S.registry()["languages"]}.get(code)
    if reg:
        for k in ("tag", "dir", "gloss"):
            if settings.get(k) and settings[k] != reg[k]: warn(f"Settings '{k}' = {settings[k]!r} but the registry says {reg[k]!r}")
    else: warn(f"'{code}' is not in tools/languages_registry.json (add it so the flag and sorting are known)")
    if missing: warn(f"{len(missing)} cell(s) left blank, English used instead: {', '.join(missing[:12])}{' ...' if len(missing) > 12 else ''}")
    if same: warn(f"{len(same)} cell(s) identical to English (fine for names/numbers, otherwise untranslated): {', '.join(same[:12])}{' ...' if len(same) > 12 else ''}")
    return pack, len(S.FIELDS) + len(S.SECTION_IDS) + len(S.QUESTION_IDS) - len(missing)

def write_pack(repo, code, pack):
    path = os.path.join(repo, "lang", f"{code}.js")
    open(path, "w", encoding="utf-8").write(f"/* Language pack: {pack['name']} ({code}). Loaded on demand by index.html. */\nwindow.registerLanguage({json.dumps(pack, ensure_ascii=False, indent=2)});\n")
    return path

# ------------------------------------------------------------------ 2. audio
def probe(p):
    o = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "stream=channels:format=duration,bit_rate", "-of", "json", p], capture_output=True, text=True)
    j = json.loads(o.stdout or "{}")
    return int(j["streams"][0]["channels"]), float(j["format"]["duration"]), int(j["format"].get("bit_rate", 0))

def process_audio(src, code, pack, repo):
    tmp = None
    if src.lower().endswith(".zip"):
        tmp = tempfile.mkdtemp(); zipfile.ZipFile(src).extractall(tmp); src = tmp
    found, bad_prefix, odd = {}, [], []          # bad_prefix: (file name, slot) pairs
    for dp, _, fs in os.walk(src):
        if "__MACOSX" in dp: continue
        for f in fs:
            if f.startswith("._") or not f.lower().endswith((".mp3", ".wav", ".m4a")): continue
            m = AUDIO_RE.match(f)
            if not m: odd.append(f); continue
            if m["prefix"].lower() != code.lower(): bad_prefix.append((f, m["slot"].lower())); continue
            slot = m["slot"].lower()
            if slot in found: err(f"two files for {slot}: {os.path.basename(found[slot])} and {f}")
            found[slot] = os.path.join(dp, f)
    # another language's file is harmless if the correct file for that slot is here (a leftover); it is an ERROR if it
    # could be a mislabelled file standing in for a slot we don't have
    stray = [f for f, slot in bad_prefix if slot in found]; risky = [f for f, slot in bad_prefix if slot not in found]
    if risky: err(f"{len(risky)} audio file(s) carry another language's prefix (not '{code}_') and their slot has no '{code}_' file: {', '.join(risky[:6])}. Rename them if they really are {code}, otherwise remove them.")
    if stray: warn(f"{len(stray)} leftover file(s) from another language ignored: {', '.join(stray[:6])}")
    if odd: warn(f"{len(odd)} file(s) ignored because the name doesn't fit <code>_intro / _s1..s6 / _q01..q15 .mp3: {', '.join(odd[:6])}")
    outdir = os.path.join(repo, "audio", code); os.makedirs(outdir, exist_ok=True)
    stats, written = {}, []
    for slot, p in sorted(found.items()):
        ch, dur, br = probe(p); dest = os.path.join(outdir, slot + ".mp3")
        if not p.lower().endswith(".mp3") or ch != 1 or br > 110000:
            subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", p, "-b:a", "96k", "-ac", "1", dest], check=True); how = "re-encoded"
        else: shutil.copyfile(p, dest); how = "kept"
        stats[slot] = {"dur": probe(dest)[1], "how": how, "kb": os.path.getsize(dest) // 1024}; written.append(dest)
        if os.path.getsize(dest) > 2 * 1024 * 1024 and slot != "intro": warn(f"{slot}.mp3 is {os.path.getsize(dest)//1024} KB: unusually large")
    # --- length sanity: does each question clip's length fit its text?
    qs = [(i, stats[f"q{i+1:02d}"]["dur"]) for i in range(15) if f"q{i+1:02d}" in stats]
    corr = None
    if len(qs) >= 8:
        xs = [len(re.sub(r"\s", "", pack["questions"][i]["t"] + pack["questions"][i]["q"])) for i, _ in qs]; ys = [d for _, d in qs]
        mx, my = statistics.mean(xs), statistics.mean(ys); sxx = sum((x - mx) ** 2 for x in xs); syy = sum((y - my) ** 2 for y in ys)
        corr = (sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / (sxx * syy) ** 0.5) if sxx and syy else 0
        # robust line (Theil-Sen): a couple of wrong clips must not hide themselves by inflating the tolerance
        slopes = [(ys[j] - ys[i]) / (xs[j] - xs[i]) for i in range(len(xs)) for j in range(i + 1, len(xs)) if xs[j] != xs[i]]
        if slopes:
            b = statistics.median(slopes); a = statistics.median([y - b * x for x, y in zip(xs, ys)])
            res = [y - (a + b * x) for x, y in zip(xs, ys)]; med = statistics.median(res)
            sd = 1.4826 * statistics.median([abs(e - med) for e in res])
            odd_q = [f"q{i+1:02d} ({d:.1f}s, text suggests ~{a+b*x:.1f}s)" for (i, d), x, e in zip(qs, xs, res) if abs(e) > max(3 * sd, 1.5)]
            if odd_q: warn("check these question clips: their length doesn't fit their text (swapped or wrong recording?): " + "; ".join(odd_q))
        if corr < 0.8: warn(f"question clip lengths barely track their text (r={corr:.2f}): files may be in the wrong order")
    for slot, st_ in stats.items():
        lo, hi = (20, 240) if slot == "intro" else ((0.6, 9) if slot.startswith("s") else (3, 40))
        if not lo <= st_["dur"] <= hi: warn(f"{slot}.mp3 is {st_['dur']:.1f}s: outside the usual {lo}-{hi}s")
    return stats, written, corr, tmp

# ------------------------------------------------------------------ 3. flag + languages.js
def ensure_flag(repo, code, settings):
    p = os.path.join(repo, "flags", code + ".svg")
    if os.path.exists(p): return True
    cc = settings.get("flag") or ({l["code"]: l for l in S.registry()["languages"]}.get(code) or {}).get("flag")
    if not cc: err(f"flags/{code}.svg is missing and no flag country is known"); return False
    try:
        d = tempfile.mkdtemp(); subprocess.run(["npm", "install", "--prefix", d, "flag-icons", "--silent"], check=True, capture_output=True)
        svg = open(os.path.join(d, "node_modules", "flag-icons", "flags", "4x3", cc + ".svg"), encoding="utf-8").read()
        if "xlink:" in svg and "xmlns:xlink" not in svg: svg = svg.replace("<svg ", '<svg xmlns:xlink="http://www.w3.org/1999/xlink" ', 1)
        open(p, "w", encoding="utf-8").write(svg); notes.append(f"flag fetched for '{cc}'"); return True
    except Exception as e:
        err(f"flags/{code}.svg is missing and could not be fetched ({e})"); return False

def update_languages_js(repo, code, name, gloss, has_audio):
    p = os.path.join(repo, "languages.js"); s = open(p, encoding="utf-8").read()
    pat = re.compile(r"\{\s*code:\s*\"([^\"]+)\"\s*,\s*name:\s*\"([^\"]*)\"\s*,\s*gloss:\s*\"([^\"]*)\"\s*,\s*audio:\s*(true|false)\s*\}")
    block = re.search(r"window\.LANGUAGES\s*=\s*\[(.*?)\];", s, re.S)
    entries = [(m[1], m[2], m[3], m[4] == "true") for m in pat.finditer(block.group(1))]
    entries = [e for e in entries if e[0] != code] + [(code, name, gloss, has_audio)]
    J = lambda x: json.dumps(x, ensure_ascii=False)
    new = "window.LANGUAGES = [\n" + ",\n".join(f"  {{ code: {J(c)}, name: {J(n)}, gloss: {J(g)}, audio: {'true' if a else 'false'} }}" for c, n, g, a in entries) + "\n];"
    s = s[:block.start()] + new + s[block.end():]
    old = re.search(r'window\.ASSET_VERSION\s*=\s*"([^"]*)"', s).group(1); today = datetime.date.today().isoformat()
    m = re.match(r"^(\d{4}-\d{2}-\d{2})-(\d+)$", old); date, n = (m[1], int(m[2])) if m else (today, 0)
    if date < today: date, n = today, 0
    newv = f"{date}-{n+1}"; s = re.sub(r'(window\.ASSET_VERSION\s*=\s*)"[^"]*"', rf'\g<1>"{newv}"', s)
    open(p, "w", encoding="utf-8").write(s); return len(entries), old, newv

# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", required=True); ap.add_argument("--audio"); ap.add_argument("--repo", default=S.ROOT); ap.add_argument("--out", default=None)
    a = ap.parse_args(); repo = os.path.abspath(a.repo); out = a.out or os.path.join(repo, "dist"); os.makedirs(out, exist_ok=True)
    settings, tr, en_tpl = read_template(a.template); code = (settings.get("code") or "").strip().lower()
    if not CODE_RE.match(code): sys.exit(f"Settings: language code '{code}' is not valid (use e.g. fr or pt-br)")
    pack, nfilled = build_pack(code, settings, tr, en_tpl, repo)
    if errors: print("ERRORS:\n  " + "\n  ".join(errors)); print("Stopped: fix the spreadsheet and run again."); sys.exit(1)
    write_pack(repo, code, pack)
    stats, written, corr, tmp = ({}, [], None, None)
    if a.audio: stats, written, corr, tmp = process_audio(a.audio, code, pack, repo)
    existing = [f for f in os.listdir(os.path.join(repo, "audio", code))] if os.path.isdir(os.path.join(repo, "audio", code)) else []
    has_audio = bool(existing)
    flag_ok = ensure_flag(repo, code, settings)
    n_lang, v_old, v_new = update_languages_js(repo, code, pack["name"], pack["gloss"], has_audio)
    chk = subprocess.run([sys.executable, os.path.join(repo, "tools", "check_languages.py")], capture_output=True, text=True, cwd=repo)
    chk_tail = chk.stdout.strip().splitlines()[-1] if chk.stdout.strip() else chk.stderr.strip()[-200:]
    if chk.returncode != 0: err("checker failed: " + chk_tail)
    # --- package
    files = ["languages.js", f"lang/{code}.js"] + ([f"flags/{code}.svg"] if flag_ok else []) + [os.path.relpath(w, repo) for w in written]
    zpath = os.path.join(out, f"{code}-update.zip")
    if not errors:                                   # never hand over a package that has a known problem
        with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
            for f in files: z.write(os.path.join(repo, f), f)
    elif os.path.exists(zpath): os.remove(zpath)
    if tmp: shutil.rmtree(tmp, ignore_errors=True)
    # --- report
    print(f"== {code}  {pack['name']} / {pack['gloss']}  (tag {pack['tag']}, {pack['dir']}) ==")
    print(f"TEXT   : {nfilled}/{len(S.FIELDS)+6+15} cells translated; pack written to lang/{code}.js")
    if a.audio:
        have = sorted(stats); miss = [s for s in SLOTS if s not in stats]
        print(f"AUDIO  : {len(have)}/22 clips this run ({sum(1 for v in stats.values() if v['how']=='re-encoded')} re-encoded to 96 kbps mono, {sum(1 for v in stats.values() if v['how']=='kept')} already fine); {sum(v['kb'] for v in stats.values())//1024 or 1} MB"
              + (f"; length-vs-text r={corr:.2f}" if corr is not None else "") + (f"; not in this batch: {', '.join(miss)}" if miss else ""))
    else: print("AUDIO  : none supplied (text-only: tile will use the phone's built-in voice)")
    print(f"FLAG   : {'flags/'+code+'.svg ok' if flag_ok else 'MISSING'}")
    print(f"LIST   : languages.js now lists {n_lang} languages; ASSET_VERSION {v_old} -> {v_new}")
    print(f"CHECK  : {chk_tail}")
    for w in warnings: print("WARN   :", w)
    for e in errors: print("ERROR  :", e)
    for n in notes: print("NOTE   :", n)
    if warnings and not errors: print(f"REVIEW : {len(warnings)} warning(s) above need a decision before publishing")
    if errors: print("STOPPED : fix the ERROR line(s) above and run again. No package was made.")
    else: print(f"PACKAGE: {zpath}  ({len(files)} files)  <- upload exactly this")
    sys.exit(1 if errors else 0)

if __name__ == "__main__": main()
