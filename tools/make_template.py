#!/usr/bin/env python3
"""
Builds the translator spreadsheet(s). English is pulled from lang/en.js, so templates are always current.

  python3 tools/make_template.py --blank  --out dist/templates      # one generic blank template
  python3 tools/make_template.py --code fr --out dist/templates      # one language, settings pre-filled from the registry
  python3 tools/make_template.py --all    --out dist/templates      # generic + every 'planned'/'pending' language

Needs:  pip install openpyxl
"""
import argparse, math, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lang_schema as S

GREY = PatternFill("solid", fgColor="EDEDED"); INPUT = PatternFill("solid", fgColor="FFF8DC")
HEAD = PatternFill("solid", fgColor="3B4F3A"); DONE = PatternFill("solid", fgColor="DDF0DD")
WRAP = Alignment(wrap_text=True, vertical="top"); thin = Side(style="thin", color="CCCCCC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HEAD; cell.alignment = WRAP; cell.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

def style(cell, kind):
    cell.alignment = WRAP; cell.border = BOX
    if kind == "en": cell.fill = GREY; cell.font = Font(color="555555")
    if kind == "in": cell.fill = INPUT
    if kind == "id": cell.font = Font(color="999999", size=8)

def fit_row(ws, r, texts, widths_chars, factor=1.35):
    lines = max(math.ceil(len(t or "") * factor / w) for t, w in zip(texts, widths_chars))
    ws.row_dimensions[r].height = max(18, min(300, 15.5 * max(1, lines)))

def greenify(ws, rng, first_cell):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"LEN({first_cell})>0"], fill=DONE))

def build(path, reg=None):
    en = S.load_pack(os.path.join(S.ROOT, "lang", "en.js"))
    wb = Workbook()
    # ---------------------------------------------------------------- READ ME (filled last)
    rd = wb.active; rd.title = "READ ME"
    # ---------------------------------------------------------------- Settings
    st = wb.create_sheet("Settings"); header(st, ["id", "Setting", "Value (fill in the yellow cells)", "Notes"], [8, 32, 40, 70])
    for r, (sid, label, note) in enumerate(S.SETTINGS, 2):
        val = (reg or {}).get(sid if sid != "note" else "note", "")
        st.cell(r, 1, sid); st.cell(r, 2, label); st.cell(r, 3, val); st.cell(r, 4, note)
        style(st.cell(r, 1), "id"); style(st.cell(r, 2), "en"); style(st.cell(r, 3), "in"); style(st.cell(r, 4), "en")
    dv = DataValidation(type="list", formula1='"ltr,rtl"', allow_blank=True); st.add_data_validation(dv); dv.add("C6")
    # ---------------------------------------------------------------- Page 2 / Buttons & labels
    rows = {}   # id -> (sheet, row)
    for sheet_name in ("Page 2", "Buttons & labels"):
        ws = wb.create_sheet(sheet_name)
        header(ws, ["id", "Where it appears", "English (do not edit)", "Your translation"], [8, 40, 60, 60])
        r = 2
        for sh, fid, where, fpath, _ in S.FIELDS:
            if sh != sheet_name: continue
            english = S.get_path(en, fpath)
            ws.cell(r, 1, fid); ws.cell(r, 2, where); ws.cell(r, 3, english); ws.cell(r, 4, None)
            style(ws.cell(r, 1), "id"); style(ws.cell(r, 2), "en"); style(ws.cell(r, 3), "en"); style(ws.cell(r, 4), "in")
            fit_row(ws, r, [where, english], [40, 60]); rows[fid] = (sheet_name, r); r += 1
        greenify(ws, f"D2:D{r-1}", "D2")
    # ---------------------------------------------------------------- Sections
    sc = wb.create_sheet("Sections"); header(sc, ["id", "Where it appears", "English (do not edit)", "Your translation"], [8, 40, 60, 60])
    for i, sid in enumerate(S.SECTION_IDS):
        r = i + 2
        sc.cell(r, 1, sid); sc.cell(r, 2, f"Section {i+1} heading (also read aloud as one recording)"); sc.cell(r, 3, en["sections"][i]); sc.cell(r, 4, None)
        style(sc.cell(r, 1), "id"); style(sc.cell(r, 2), "en"); style(sc.cell(r, 3), "en"); style(sc.cell(r, 4), "in")
        fit_row(sc, r, [en["sections"][i]], [60]); rows[sid] = ("Sections", r)
    greenify(sc, f"D2:D{len(S.SECTION_IDS)+1}", "D2")
    # ---------------------------------------------------------------- Questions
    qs = wb.create_sheet("Questions")
    header(qs, ["id", "#", "English title (do not edit)", "English question (do not edit)", "Title (your translation)", "Question (your translation)"], [6, 4, 28, 58, 30, 62])
    for i, qid in enumerate(S.QUESTION_IDS):
        r = i + 2; q = en["questions"][i]
        for c, v in enumerate([qid, i + 1, q["t"], q["q"], None, None], 1): qs.cell(r, c, v)
        style(qs.cell(r, 1), "id"); style(qs.cell(r, 2), "en"); style(qs.cell(r, 3), "en"); style(qs.cell(r, 4), "en")
        style(qs.cell(r, 5), "in"); style(qs.cell(r, 6), "in"); fit_row(qs, r, [q["q"]], [58]); rows[qid] = ("Questions", r)
    greenify(qs, f"E2:F{len(S.QUESTION_IDS)+1}", "E2")
    # ---------------------------------------------------------------- Audio (what to record, file names)
    au = wb.create_sheet("Audio"); header(au, ["File name (use exactly)", "What it is", "Script to read aloud (from your translation)", "Recorded? (Y/N)"], [26, 34, 90, 14])
    code = "Settings!$C$2"
    def ref(fid): sh, r = rows[fid]; return f"'{sh}'!D{r}"
    intro_parts = [ref("introBody"), ref("how1"), ref("how2"), ref("how3")]
    intro_parts += [f'{ref(f"legend{i}_label")}&" — "&{ref(f"legend{i}_desc")}' for i in range(1, 5)]
    intro_parts += [ref("beforeYouBegin"), ref("formalDisclaimer")]
    audio_rows = [(f'={code}&"_intro.mp3"', "Page 2: read ALL of it, one line after the other (about 1-2 minutes)", "=" + '&CHAR(10)&'.join(intro_parts))]
    for i in range(N := S.N_SECTIONS):
        audio_rows.append((f'={code}&"_s{i+1}.mp3"', f"Section {i+1} heading (just the section name)", f"=Sections!D{i+2}"))
    for i in range(S.N_QUESTIONS):
        audio_rows.append((f'={code}&"_q{i+1:02d}.mp3"', f"Question {i+1}: the title, then the question", f'=Questions!E{i+2}&". "&Questions!F{i+2}'))
    for r, (f, what, script) in enumerate(audio_rows, 2):
        au.cell(r, 1, f); au.cell(r, 2, what); au.cell(r, 3, script); au.cell(r, 4, None)
        style(au.cell(r, 1), "en"); style(au.cell(r, 2), "en"); style(au.cell(r, 3), "en"); style(au.cell(r, 4), "in")
        au.row_dimensions[r].height = 190 if r == 2 else (30 if r < 9 else 48)
    dv2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); au.add_data_validation(dv2); dv2.add(f"D2:D{len(audio_rows)+1}")
    # ---------------------------------------------------------------- READ ME
    rd.column_dimensions["A"].width = 3; rd.column_dimensions["B"].width = 120
    n_p2 = sum(1 for f in S.FIELDS if f[0] == "Page 2"); n_bt = sum(1 for f in S.FIELDS if f[0] == "Buttons & labels")
    text = [
     ("The Menopause Conversation: language template", "title"),
     ("This spreadsheet carries everything for ONE new language: the words people read, and the list of voice recordings to make.", None),
     ("", None),
     ("FOR THE TRANSLATOR", "h"),
     ("1.  Fill the yellow cells on the 'Settings' sheet (or check them if already filled).", None),
     ("2.  On each sheet, translate the grey English into the yellow cell next to it. Never edit the grey cells. A cell turns green once it has text.", None),
     ("3.  Keep {n} exactly as it is wherever it appears. The app swaps in a number there.", None),
     ("4.  Do not add numbers like '1.' to the steps or questions. The app numbers them itself.", None),
     ("5.  Use one consistent politeness level throughout (formal or friendly, but not a mix).", None),
     ("6.  If a phrase has no natural translation, write the closest natural wording and tell us in the 'Variant / script note' setting.", None),
     ("7.  Medical wording should be read by a native speaker before this is sent back. Put their name in 'Reviewed by'.", None),
     ("", None),
     ("FOR THE PERSON RECORDING", "h"),
     ("1.  The 'Audio' sheet lists all 22 recordings, the exact file name for each, and the exact script to read (it fills in from the translation).", None),
     ("2.  Save every file with EXACTLY that name, e.g. ko_q01.mp3. The language code in front is what stops different languages being mixed up.", None),
     ("3.  Export settings: MP3, mono (1 channel), 96 kbps. Smaller files load faster on phones.", None),
     ("4.  Put all 22 files in one folder and zip it as <code>_audio.zip. Audio can arrive after the text; send the zip again when you have more.", None),
     ("", None),
     ("PROGRESS (updates as you type)", "h"),
     (f'="Page 2:  "&COUNTA(\'Page 2\'!D2:D{n_p2+1})&" of {n_p2}"', None),
     (f'="Buttons & labels:  "&COUNTA(\'Buttons & labels\'!D2:D{n_bt+1})&" of {n_bt}"', None),
     (f'="Section names:  "&COUNTA(Sections!D2:D{S.N_SECTIONS+1})&" of {S.N_SECTIONS}"', None),
     (f'="Questions (titles + questions):  "&(COUNTA(Questions!E2:E{S.N_QUESTIONS+1})+COUNTA(Questions!F2:F{S.N_QUESTIONS+1}))&" of {2*S.N_QUESTIONS}"', None),
     (f'="Recordings marked done:  "&COUNTIF(Audio!D2:D{len(audio_rows)+1},"Y")&" of {len(audio_rows)}"', None),
    ]
    for i, (t, kind) in enumerate(text, 1):
        c = rd.cell(i, 2, t); c.alignment = WRAP
        if kind == "title": c.font = Font(bold=True, size=16, color="3B4F3A")
        if kind == "h": c.font = Font(bold=True, size=12, color="954A63")
    for ws in wb.worksheets:                      # print-friendly: landscape, one page wide
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.save(path)
    return path

def slug(s): return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="dist/templates"); g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--blank", action="store_true"); g.add_argument("--code"); g.add_argument("--all", action="store_true")
    a = ap.parse_args(); os.makedirs(a.out, exist_ok=True); made = []
    reg = {l["code"]: l for l in S.registry()["languages"]}
    if a.blank or a.all: made.append(build(os.path.join(a.out, "Language_Template_BLANK.xlsx"), {"dir": "ltr"}))
    codes = [a.code] if a.code else ([c for c, l in reg.items() if l["status"] in ("planned", "pending")] if a.all else [])
    for c in codes:
        if c not in reg: sys.exit(f"'{c}' is not in tools/languages_registry.json. Add it there first.")
        made.append(build(os.path.join(a.out, f"{c}_{slug(reg[c]['gloss'])}_template.xlsx"), reg[c]))
    print("\n".join(made)); print(f"{len(made)} template(s) written")

if __name__ == "__main__": main()
